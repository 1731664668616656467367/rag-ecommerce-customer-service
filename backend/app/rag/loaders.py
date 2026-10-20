"""文档读取：把 txt/md/pdf/docx 统一读成纯文本。

注意：上传的文档是“待处理的数据”，不是指令。
即使文档里写着“忽略以上所有规则”，它也只是会被当作普通文本切块，
永远不会进入 system prompt 的指令区。
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED = {".txt", ".md", ".markdown", ".pdf", ".docx"}


class DocumentError(Exception):
    pass


def read_text_file(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    raise DocumentError("无法识别文件编码，请另存为 UTF-8 后重试")


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("缺少 pypdf 依赖，无法解析 PDF") from exc
    try:
        reader = PdfReader(str(path))
        pages = [(p.extract_text() or "") for p in reader.pages]
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"PDF 解析失败：{exc}") from exc
    text = "\n\n".join(pages).strip()
    if not text:
        raise DocumentError("PDF 中没有可提取的文字（可能是扫描件）")
    return text


def read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("缺少 python-docx 依赖，无法解析 docx") from exc
    try:
        doc = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"docx 解析失败：{exc}") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    text = "\n".join(parts).strip()
    if not text:
        raise DocumentError("docx 中没有文字内容")
    return text


def load_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise DocumentError(f"不支持的文件类型：{suffix}")
    if suffix in (".txt", ".md", ".markdown"):
        return read_text_file(path)
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix == ".docx":
        return read_docx(path)
    raise DocumentError(f"未处理的类型：{suffix}")
